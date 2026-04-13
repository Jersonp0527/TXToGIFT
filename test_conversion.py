#!/usr/bin/env python3
"""
Prueba de extremo a extremo: genera un .docx de ejemplo con los
tres tipos de pregunta y verifica que la conversión GIFT sea correcta.
"""

import sys
import tempfile
import os
from docx import Document
from docx_to_gift import parse_docx, convert_to_gift


def build_sample_docx(path: str):
    doc = Document()

    # 1. Selección múltiple
    doc.add_paragraph(
        "De los siguientes aspectos de la globalización, cuál es FALSO:"
    )
    doc.add_paragraph(
        "A.\tSe refiere al cambio hacia una economía mundial más integrada e interdependiente."
    )
    doc.add_paragraph(
        "B.\tEs un proceso histórico, el resultado de la innovación humana y el progreso tecnológico."
    )
    doc.add_paragraph(
        "C.\tSus dos principales impulsores son: la reducción de las barreras al comercio a las inversiones, y la integración económica regional."
    )
    doc.add_paragraph(
        "D.\tAcuerdos como el GATT han mejorado la integración de los mercados y la interdependencia a nivel global."
    )
    doc.add_paragraph("ANSWER: C")

    # 2. Emparejamiento (matching)
    doc.add_paragraph(
        "Por favor empareje los siguientes elementos del costo total:"
    )
    doc.add_paragraph(
        "A.\tInventarios en planta, manejo en planta → Costo de calidad y operaciones"
    )
    doc.add_paragraph(
        "B.\tCosto de riesgo, personal de compras → Otros costos"
    )
    doc.add_paragraph(
        "C.\tCostos de materiales directos, mano de obra directa → Precio del proveedor"
    )
    doc.add_paragraph(
        "D.\tTransporte en el país, fletes marítimos/aéreos → Costos de entrega"
    )
    doc.add_paragraph("ANSWER: ")

    # 3. Verdadero / Falso  (respuesta verdadera)
    doc.add_paragraph(
        "La ventaja principal del franquiciamiento consiste en que el comprador asume los costos y riesgos."
    )
    doc.add_paragraph("A.\tVerdadero")
    doc.add_paragraph("B.\tFalso")
    doc.add_paragraph("ANSWER: A")

    # 4. Verdadero / Falso  (respuesta falsa)
    doc.add_paragraph(
        "El licenciamiento consiste en la venta de una propiedad intangible al licenciante."
    )
    doc.add_paragraph("A.\tVerdadero")
    doc.add_paragraph("B.\tFalso")
    doc.add_paragraph("ANSWER: B")

    doc.save(path)


def test():
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tf:
        docx_path = tf.name

    try:
        build_sample_docx(docx_path)
        questions = parse_docx(docx_path)
        gift = convert_to_gift(questions)

        print("=== GIFT generado ===")
        print(gift)

        # Verificaciones básicas
        assert len(questions) == 4, f"Se esperaban 4 preguntas, se obtuvieron {len(questions)}"
        assert '{TRUE}' in gift,  "Falta {TRUE} en la salida"
        assert '{FALSE}' in gift, "Falta {FALSE} en la salida"
        assert '=Sus dos principales' in gift, "La opción correcta de la P1 no tiene '='"
        assert '~Se refiere' in gift,          "Una opción incorrecta de la P1 no tiene '~'"
        assert '=Inventarios en planta' in gift, "La P2 (matching) no contiene '=Inventarios'"
        assert '->' in gift,                     "La P2 (matching) no contiene '->'"

        print("Todas las verificaciones pasaron correctamente.")
        return True

    except AssertionError as e:
        print(f"FALLO: {e}", file=sys.stderr)
        return False
    finally:
        os.unlink(docx_path)


if __name__ == '__main__':
    ok = test()
    sys.exit(0 if ok else 1)
